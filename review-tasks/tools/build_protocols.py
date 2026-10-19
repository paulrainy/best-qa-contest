"""Протоколы НТ задания 1г: Markdown-исходник → .docx.

Запуск: python build_protocols.py ../1g-load-review
Читает jury/source/protocols/NN-*.md, пишет participant/protocols/NN-*.docx.

Исходник:
- строки «> …» до названия — шапка (выравнивается вправо), «# …» — название;
- «## N. …» — раздел, «### N.M …» — подраздел, номера идут подряд;
- «Таблица N. …» — подпись, сразу за ней таблица Markdown;
- «![Рисунок N. …](img/файл.png)» — рисунок, путь от папки исходника;
- «- …» — пункт списка, «**…**» — жирный текст, остальное — абзацы.
Падает, если номера идут не подряд, нет картинки, у таблицы нет подписи или в строке не то число колонок.
В свойствах документа (identifier) — хэш исходника и картинок: по нему build_sheets видит устаревший документ.
"""
import hashlib
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = 'Times New Roman'
IMAGE = re.compile(r'^!\[(Рисунок (\d+)\.[^\]]*)\]\(([^)]+)\)\s*$')
CAPTION = re.compile(r'^Таблица (\d+)\.')
SECTION = re.compile(r'^## (\d+)\. ')
SUBSECTION = re.compile(r'^### (\d+)\.(\d+) ')
BOLD = re.compile(r'\*\*(.+?)\*\*')


@dataclass
class Block:
    kind: str  # heading, para, list, table, figure
    text: str = ''
    level: int = 0
    rows: list = field(default_factory=list)
    image: Path | None = None
    number: str = ''


@dataclass
class Protocol:
    header: list
    title: str
    blocks: list


def cells(line: str) -> list:
    return [c.strip() for c in line.strip().strip('|').split('|')]


def parse(path: Path) -> Protocol:
    lines = path.read_text(encoding='utf-8').splitlines()
    header, title, blocks, errors = [], '', [], []
    last = {'Таблица': None, 'Рисунок': None}
    section, subsection = 0, 0

    def follows(what, number):
        before = last[what]
        if number != (before or 0) + 1:
            errors.append(f'{what} {number} после ' + (f'{what} {before}' if before else 'начала'))
        last[what] = number

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.startswith('> ') and not title:
            header.append(line[2:].strip())
        elif line.startswith('# '):
            title = line[2:].strip()
        elif SECTION.match(line):
            number = int(SECTION.match(line).group(1))
            if number != section + 1:
                errors.append(f'раздел {number} после ' + (str(section) if section else 'начала'))
            section, subsection = number, 0
            blocks.append(Block('heading', line[3:].strip(), 2, number=str(number)))
        elif SUBSECTION.match(line):
            major, minor = map(int, SUBSECTION.match(line).groups())
            if major != section or minor != subsection + 1:
                errors.append(f'раздел {major}.{minor} после {section}' + (f'.{subsection}' if subsection else ''))
            subsection = minor
            blocks.append(Block('heading', line[4:].strip(), 3, number=f'{major}.{minor}'))
        elif IMAGE.match(line):
            caption, number, file = IMAGE.match(line).groups()
            follows('Рисунок', int(number))
            image = path.parent / file
            if not image.exists():
                errors.append(f'нет файла {file}')
            blocks.append(Block('figure', caption, image=image, number=number))
        elif CAPTION.match(line):
            number = int(CAPTION.match(line).group(1))
            follows('Таблица', number)
            j = i + 1
            while j < len(lines) and not lines[j].strip():
                j += 1
            table = []
            while j < len(lines) and lines[j].strip().startswith('|'):
                table.append(lines[j])
                j += 1
            if len(table) < 3:
                errors.append(f'Таблица {number}: нет таблицы под подписью')
            else:
                head, rows = cells(table[0]), [cells(r) for r in table[2:]]
                for n, row in enumerate(rows, start=1):
                    if len(row) != len(head):
                        errors.append(f'Таблица {number}: строка {n} — {len(row)} колонки, нужно {len(head)}')
                blocks.append(Block('table', line.strip(), rows=[head, *rows], number=str(number)))
            i = j
            continue
        elif line.lstrip().startswith('|'):
            errors.append(f'таблица без подписи (строка {i + 1})')
            while i < len(lines) and lines[i].strip().startswith('|'):
                i += 1
            continue
        elif line.startswith('- '):
            blocks.append(Block('list', line[2:].strip()))
        else:
            if blocks and blocks[-1].kind == 'para' and i > 0 and lines[i - 1].strip():
                blocks[-1].text += ' ' + line.strip()
            else:
                blocks.append(Block('para', line.strip()))
        i += 1
    if not title:
        errors.append('нет названия «# …»')
    if errors:
        sys.exit(f'{path.name}: ' + '; '.join(errors))
    return Protocol(header, title, blocks)


def places(protocol: Protocol) -> dict:
    found = {'шапка': ' / '.join(protocol.header), 'название': protocol.title}
    for block in protocol.blocks:
        if block.kind == 'heading':
            found[f'раздел:{block.number}'] = block.text
        elif block.kind == 'table':
            found[f'таблица:{block.number}'] = block.text
        elif block.kind == 'figure':
            found[f'рисунок:{block.number}'] = block.text
    if not protocol.header:
        del found['шапка']
    return found


def source_hash(path: Path) -> str:
    digest = hashlib.sha256(path.read_bytes())
    for line in path.read_text(encoding='utf-8').splitlines():
        match = IMAGE.match(line.rstrip())
        if match and (path.parent / match.group(3)).exists():
            digest.update((path.parent / match.group(3)).read_bytes())
    return digest.hexdigest()[:16]


def runs(paragraph, text: str, bold=False, italic=False):
    for n, part in enumerate(BOLD.split(text)):
        if part:
            run = paragraph.add_run(part)
            run.bold = bold or n % 2 == 1 or None
            run.italic = italic or None


def style_fonts(doc):
    for name in ('Normal', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet', 'Title'):
        style = doc.styles[name]
        style.font.name = FONT
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), FONT)
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles['Normal'].font.size = Pt(12)
    for name, size in (('Heading 1', 14), ('Heading 2', 13), ('Heading 3', 12)):
        doc.styles[name].font.size = Pt(size)


def write(protocol: Protocol, md_path: Path, out: Path) -> None:
    doc = Document()
    style_fonts(doc)
    for line in protocol.header:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        runs(paragraph, line)
    title = doc.add_heading(level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    runs(title, protocol.title)
    for block in protocol.blocks:
        if block.kind == 'heading':
            runs(doc.add_heading(level=block.level), block.text)
        elif block.kind == 'para':
            runs(doc.add_paragraph(), block.text)
        elif block.kind == 'list':
            runs(doc.add_paragraph(style='List Bullet'), block.text)
        elif block.kind == 'table':
            runs(doc.add_paragraph(), block.text, italic=True)
            head, *rows = block.rows
            table = doc.add_table(rows=1 + len(rows), cols=len(head))
            table.style = 'Table Grid'
            for r, row in enumerate([head, *rows]):
                for c, value in enumerate(row):
                    cell = table.cell(r, c).paragraphs[0]
                    runs(cell, value, bold=r == 0)
            doc.add_paragraph()
        elif block.kind == 'figure':
            doc.add_picture(str(block.image), width=Cm(16))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption = doc.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            runs(caption, block.text, italic=True)
    doc.core_properties.title = protocol.title
    doc.core_properties.author = 'ЦК НТ'
    doc.core_properties.identifier = source_hash(md_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)


def build(task: Path) -> list:
    built = []
    for md in sorted((task / 'jury' / 'source' / 'protocols').glob('[0-9][0-9]-*.md')):
        out = task / 'participant' / 'protocols' / f'{md.stem}.docx'
        write(parse(md), md, out)
        built.append(out)
    return built


if __name__ == '__main__':
    if '--help' in sys.argv or len(sys.argv) < 2:
        sys.exit(__doc__)
    done = build(Path(sys.argv[1]).resolve())
    print('Готово:', ', '.join(p.name for p in done))
