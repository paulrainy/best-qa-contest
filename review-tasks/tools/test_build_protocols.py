from pathlib import Path

import matplotlib

matplotlib.use('Agg')
import matplotlib.pyplot as plt
import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

import build_protocols as bp

SOURCE = '''> УТВЕРЖДАЮ
> Руководитель ЦК НТ
# Протокол нагрузочного тестирования системы ЕФР

## 1. Сроки
С 01.11.2026 по 12.11.2026.

## 2. Результаты
Текст с **жирным** словом.

### 2.1 Поиск максимума
Таблица 1. Бизнес-метрики
| Операция | p95, мс | Порог, мс | Ошибки | Итог |
| --- | --- | --- | --- | --- |
| ОП-01 | 640 | 800 | 0,1 % | **ок** |
| ОП-02 | 910 | 1 000 | 0,2 % | ок |

Таблица 2. Утилизация
| Сервер | CPU |
| --- | --- |
| efr-nt-app01 | 45 % |

![Рисунок 1. Интенсивность](img/x.png)

- первый пункт
- второй пункт
'''


def picture(path: Path, color='C0'):
    path.parent.mkdir(parents=True, exist_ok=True)
    fig, ax = plt.subplots(figsize=(2, 1))
    ax.plot([0, 1], [0, 1], color=color)
    fig.savefig(path)
    plt.close(fig)


def source(tmp_path: Path, text=SOURCE, image=True) -> Path:
    path = tmp_path / '11-max.md'
    path.write_text(text, encoding='utf-8')
    if image:
        picture(tmp_path / 'img' / 'x.png')
    return path


def test_parse_places(tmp_path):
    places = bp.places(bp.parse(source(tmp_path)))
    assert set(places) == {'шапка', 'название', 'раздел:1', 'раздел:2', 'раздел:2.1', 'таблица:1', 'таблица:2',
                           'рисунок:1'}
    assert places['раздел:2.1'] == '2.1 Поиск максимума'
    assert places['таблица:2'] == 'Таблица 2. Утилизация'


@pytest.mark.parametrize('old, new, image, message', [
    ('Таблица 2. Утилизация', 'Таблица 3. Утилизация', True, 'Таблица 3 после Таблица 1'),
    (None, None, False, 'нет файла img/x.png'),
    ('| ОП-02 | 910 | 1 000 | 0,2 % | ок |', '| ОП-02 | 910 | 1 000 | ок |', True,
     'Таблица 1: строка 2 — 4 колонки, нужно 5'),
    ('## 2. Результаты', '## 3. Результаты', True, 'раздел 3 после 1'),
    ('### 2.1 Поиск максимума', '### 2.2 Поиск максимума', True, 'раздел 2.2 после 2'),
    ('Таблица 2. Утилизация\n', '', True, 'таблица без подписи'),
    ('![Рисунок 1. Интенсивность]', '![Рисунок 2. Интенсивность]', True, 'Рисунок 2 после начала'),
])
def test_rejects(tmp_path, old, new, image, message):
    text = SOURCE.replace(old, new) if old else SOURCE
    with pytest.raises(SystemExit, match=message):
        bp.parse(source(tmp_path, text, image))


def test_write_docx(tmp_path):
    md = source(tmp_path)
    out = tmp_path / 'out.docx'
    bp.write(bp.parse(md), md, out)
    doc = Document(out)
    heading = next(p for p in doc.paragraphs if p.text == '2.1 Поиск максимума')
    assert heading.style.name == 'Heading 3'
    assert len(doc.tables) == 2
    assert doc.tables[0].cell(0, 1).text == 'p95, мс'
    assert len(doc.inline_shapes) == 1
    assert doc.core_properties.identifier == bp.source_hash(md)
    header = next(p for p in doc.paragraphs if p.text == 'УТВЕРЖДАЮ')
    assert header.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    para = next(p for p in doc.paragraphs if p.text == 'Текст с жирным словом.')
    assert [r.bold for r in para.runs if r.text == 'жирным'] == [True]
    cell = doc.tables[0].cell(1, 4).paragraphs[0]
    assert [r.bold for r in cell.runs] == [True]
    assert [p.style.name for p in doc.paragraphs if p.text.endswith('пункт')] == ['List Bullet', 'List Bullet']


def test_hash_changes_with_image(tmp_path):
    md = source(tmp_path)
    before = bp.source_hash(md)
    picture(tmp_path / 'img' / 'x.png', color='C3')
    assert bp.source_hash(md) != before


def test_build_writes_participant_docx(tmp_path):
    task = tmp_path / 'task'
    folder = task / 'jury' / 'source' / 'protocols'
    folder.mkdir(parents=True)
    source(folder)
    bp.build(task)
    assert (task / 'participant' / 'protocols' / '11-max.docx').exists()


def test_build_skips_up_to_date(tmp_path):
    task = tmp_path / 'task'
    folder = task / 'jury' / 'source' / 'protocols'
    folder.mkdir(parents=True)
    source(folder)
    bp.build(task)
    out = task / 'participant' / 'protocols' / '11-max.docx'
    before = out.read_bytes()
    assert bp.build(task) == []
    assert out.read_bytes() == before
