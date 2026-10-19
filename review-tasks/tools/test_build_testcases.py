from pathlib import Path

import pytest
from openpyxl import load_workbook

import build_testcases as bt

SHARED = '''
# 900901 Вход в ЕФР в роли менеджера-операциониста
Расположение: РТ_ЕФР -> Общие шаги
Автоматизирован: Нет
Приоритет: Средний
Статус: Готов

## Шаги
- Выполнить вход в ЕФР в роли менеджера-операциониста ТП
  = Открылась страница «Поиск клиента»
  > tst_oper01

# 900903 Открытие заявки контролёром
Расположение: РТ_ЕФР -> Общие шаги
Автоматизирован: Нет
Приоритет: Средний
Статус: Готов

## Шаги
- Открыть раздел «Задачи»
  = Открылся список задач
- Нажать «Взять и исполнить» у заявки
  = Открылась форма заявки
'''

EXAMPLE = '''
# 900401 Открытие вклада: сумма выше порога уходит на подтверждение контролёру
Расположение: РТ_ЕФР -> РТ_Тест-кейсы_функциональные -> BIQ-90101 Открытие вклада
Автоматизирован: Нет
Приоритет: Высокий
Статус: Готов
Дата создания: 9/14/2026 10:12:00
Автор: Ирина Соловьёва
Тег: biq90101, deposit

## Предусловия
- @900901
- Клиент с FATCA-статусом «Не является налоговым резидентом США»
  > Паспорт 4510 777001

## Шаги
- Нажать «Открыть вклад»
  = Открылся шаг «Выбор вклада»
- Ввести сумму
  = Поле «Сумма» заполнено,
    кнопка «Далее» активна
  > 1 500 000,00
- @900903

## Постусловия
- Вклад закрыт
'''


def case(case_id='900401', steps=6, biq='BIQ-90101', step_text='Шаг'):
    items = '\n'.join(f'- {step_text} {n}\n  = Результат {n}' for n in range(1, steps + 1))
    return (f'# {case_id} Кейс {case_id}\n'
            f'Расположение: РТ_ЕФР -> РТ_Тест-кейсы_функциональные -> {biq} Тема\n'
            f'Приоритет: Высокий\n\n## Шаги\n{items}\n')


def source(tmp_path: Path, text: str) -> Path:
    path = tmp_path / 'testcases.md'
    path.write_text(text, encoding='utf-8')
    return path


def rows(path: Path):
    return [list(r) for r in load_workbook(path).active.iter_rows(values_only=True)]


def test_parse_fields_and_items(tmp_path):
    cases = bt.parse(source(tmp_path, SHARED + EXAMPLE))
    example = next(c for c in cases if c.id == '900401')
    assert example.name == 'Открытие вклада: сумма выше порога уходит на подтверждение контролёру'
    assert example.fields['Приоритет'] == 'Высокий'
    assert example.steps[1].text == 'Ввести сумму'
    assert example.steps[1].expected == 'Поле «Сумма» заполнено,\nкнопка «Далее» активна'
    assert example.steps[1].data == '1 500 000,00'
    assert example.pre[1].data == 'Паспорт 4510 777001'
    assert example.post[0].text == 'Вклад закрыт'
    assert next(c for c in cases if c.id == '900903').is_shared


def test_expand_shared_steps(tmp_path):
    cases = bt.expand(bt.parse(source(tmp_path, SHARED + EXAMPLE)))
    example = next(c for c in cases if c.id == '900401')
    assert len(example.steps) == 4
    assert [s.text for s in example.steps[2:]] == ['Открыть раздел «Задачи»', 'Нажать «Взять и исполнить» у заявки']
    assert [s.shared for s in example.steps] == [None, None, '900903', '900903']
    assert example.pre[0].shared == '900901'
    assert example.pre[0].data == 'tst_oper01'


@pytest.mark.parametrize('text, kwargs, message', [
    (case() + '- @900999\n', {}, '900401: нет общего шага 900999'),
    (case(steps=5), {}, '900401: шагов 5, нужно 6–12'),
    (case(), {'ids': ['900401', '900402']}, 'нет кейса 900402'),
    (case() + case('900401'), {}, 'повтор ID 900401'),
    (case(), {'per_biq': {'BIQ-90101': 2}}, 'BIQ-90101: кейсов 1, нужно 2'),
    (case() + SHARED.replace('- Открыть раздел «Задачи»', '- @900901'), {},
     '900903: в общем шаге нельзя подключать другие'),
])
def test_check_rejects(tmp_path, text, kwargs, message):
    options = {'ids': ['900401'], 'per_biq': {}, **kwargs}
    with pytest.raises(SystemExit, match=message):
        bt.check(bt.parse(source(tmp_path, SHARED + text)), **options)


def test_check_accepts_valid(tmp_path):
    bt.check(bt.parse(source(tmp_path, SHARED + case())), ids=['900401'], per_biq={'BIQ-90101': 1})


def test_parse_rejects_unknown_line(tmp_path):
    with pytest.raises(SystemExit, match='строка 3'):
        bt.parse(source(tmp_path, '# 900401 Кейс\nРасположение: x\nчто-то непонятное\n'))


def test_write_export(tmp_path):
    out = tmp_path / 'testcases.xlsx'
    bt.write(bt.expand(bt.parse(source(tmp_path, SHARED + EXAMPLE))), out)
    table = rows(out)
    assert table[0] == bt.COLUMNS
    col = {name: i for i, name in enumerate(bt.COLUMNS)}
    header = table[1]
    assert header[col['ID']] == '900401'
    assert header[col['Наименование']].startswith('Открытие вклада')
    assert header[col['Приоритет']] == 'Высокий'
    assert header[col['№ шага']] is None
    numbers = [r[col['№ шага']] for r in table[2:9]]
    assert numbers == ['Предусловие 1', 'Предусловие 2', 'Шаг 1', 'Шаг 2', 'Шаг 3', 'Шаг 4', 'Постусловие 1']
    assert table[2][col['Предусловия']].startswith('Выполнить вход')
    assert table[3][col['Тестовые данные']] == 'Паспорт 4510 777001'
    assert table[5][col['Шаги']] == 'Ввести сумму'
    assert table[5][col['Ожидаемый результат']] == 'Поле «Сумма» заполнено,\nкнопка «Далее» активна'
    assert table[6][col['Комментарии']] == 'Общий шаг 900903 «Открытие заявки контролёром»'
    assert table[8][col['Постусловия']] == 'Вклад закрыт'
    shared_ids = [r[col['ID']] for r in table[9:] if r[col['ID']]]
    assert shared_ids == ['900901', '900903']
    sheet = load_workbook(out).active
    assert sheet.cell(row=6, column=col['Ожидаемый результат'] + 1).alignment.wrap_text


def test_numbering_restarts_per_case(tmp_path):
    out = tmp_path / 'testcases.xlsx'
    text = SHARED + case() + case('900402').replace('## Шаги', '## Предусловия\n- Клиент найден\n\n## Шаги')
    bt.write(bt.expand(bt.parse(source(tmp_path, text))), out)
    col = bt.COLUMNS.index('№ шага')
    table = rows(out)
    start = next(i for i, r in enumerate(table) if r[0] == '900402')
    assert [r[col] for r in table[start + 1:start + 3]] == ['Предусловие 1', 'Шаг 1']


def test_equals_sign_stays_text(tmp_path):
    out = tmp_path / 'testcases.xlsx'
    text = case().replace('= Результат 1', '= = 10 000,00')
    bt.write(bt.expand(bt.parse(source(tmp_path, text))), out)
    cell = load_workbook(out).active.cell(row=3, column=bt.COLUMNS.index('Ожидаемый результат') + 1)
    assert cell.data_type == 's'
    assert cell.value == '= 10 000,00'


def test_dump_numbers_items(tmp_path):
    text = bt.dump(bt.expand(bt.parse(source(tmp_path, SHARED + EXAMPLE))), '900401')
    assert 'Предусловие 2: Клиент с FATCA-статусом' in text
    assert 'Шаг 3: Открыть раздел «Задачи» [общий шаг 900903]' in text
    assert 'Шаг 3 ОР: Открылся список задач' in text
    assert 'Шаг 2 данные: 1 500 000,00' in text
