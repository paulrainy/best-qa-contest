"""Шаблон ответа участника и лист подсчёта из ключа задания.

Запуск: python build_sheets.py ../1a-autotest-review [--show]
Читает jury/answer-key.md, пишет participant/answer-template.xlsx и jury/scoring-sheet.xlsx.
Падает, если тестов не 20, в тесте не 5 дефектов, ловушек не 10 или дефект и ловушка указывают на одну строку.
--show печатает строки кода, на которые ссылается ключ, — для сверки глазами.

Задание 1б (ревью тест-кейсов): в строке дефекта есть колонка «Категория», место пишется как «Шаг 3»,
«Предусловие 1», «Наименование». Ключ сверяется с participant/testcases.xlsx: элемент должен быть в кейсе
и не должен быть взят из общего шага. --show печатает текст элементов из выгрузки.

Задание 1г (ревью материалов НТ): единица — две цифры (01–20), у дефекта категория «М» или 1–5, у ловушки
тоже есть колонка единицы. Место: путь в дереве .jmx в обратных кавычках, «Панель «…»», «Переменная host»,
«Дашборд», «строка 14», «Раздел 12.2», «Таблица 3», «Рисунок 2», «Шапка», «Название». Каждое место сверяется
с файлом единицы, сценарии и протоколы — с исходниками. --partial проверяет только места, пока ключ
пишется по частям, и листы не пишет.
"""
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

import build_testcases
import load_places

TESTS = 20
DEFECTS_PER_TEST = 5
TRAPS = 10
ROW = re.compile(r'^\|\s*(Д\d{2,3}\.\d|Л\d{1,2})\s*\|(.*)\|\s*$')
PLACE = re.compile(r'`([^`:]+):(\d+)(?:–(\d+))?`')
ELEMENT = re.compile(r'(?:(\d{6}) )?(Наименование|Предусловие \d+|Шаг \d+|Постусловие \d+)')
SECTION = re.compile(r'Предусловия|Шаги|Постусловия')
CASE_CELLS = 5  # кейс, место, категория, суть, засчитываем — у кода на одну меньше
CATEGORIES = {'М', '1', '2', '3', '4', '5'}
HEADER = Font(bold=True)
FILL = PatternFill('solid', fgColor='DDEBF7')
WRAP = Alignment(wrap_text=True, vertical='top')


def parse_key(key: Path):
    defects, traps = [], []
    for line in key.read_text(encoding='utf-8').splitlines():
        match = ROW.match(line)
        if match:
            cells = [c.strip() for c in match.group(2).split('|')]
            (defects if match.group(1).startswith('Д') else traps).append([match.group(1), *cells])
    return defects, traps


def places(text: str) -> set:
    found = set()
    for name, start, end in PLACE.findall(text):
        for line in range(int(start), int(end or start) + 1):
            found.add(f'{name}:{line}')
    return found


def kind_of(defects) -> str:
    """code — 1а и 1д, cases — 1б, load — 1г (единица из двух цифр)."""
    if defects and len(defects[0]) > 1 and re.fullmatch(r'\d{2}', defects[0][1]):
        return 'load'
    return 'cases' if defects and len(defects[0]) - 1 == CASE_CELLS else 'code'


def is_cases(defects) -> bool:
    return kind_of(defects) == 'cases'


def case_places(case: str, text: str) -> set:
    return {f'{prefix or case} {element}' for prefix, element in ELEMENT.findall(text)}


def case_elements(task: Path) -> dict:
    """Элементы выгрузки: «900401 Шаг 3» → (текст, ОР, взят ли из общего шага)."""

    path = task / 'participant' / 'testcases.xlsx'
    if not path.exists():
        return {}
    col = {name: i for i, name in enumerate(build_testcases.COLUMNS)}
    elements, case = {}, None
    for row in load_workbook(path).active.iter_rows(min_row=2, values_only=True):
        if row[col['ID']]:
            case = row[col['ID']]
            elements[f'{case} Наименование'] = (row[col['Наименование']], '', False)
        elif row[col['№ шага']]:
            text = next(row[col[c]] for c in ('Предусловия', 'Шаги', 'Постусловия') if row[col[c]] is not None)
            shared = str(row[col['Комментарии']] or '').startswith(build_testcases.SHARED_COMMENT)
            elements[f'{case} {row[col["№ шага"]]}'] = (text, row[col['Ожидаемый результат']] or '', shared)
    return elements


def source_elements(source: Path) -> dict:
    """Те же элементы, что case_elements, но прямо из исходника — чтобы поймать несобранную выгрузку."""
    elements = {}
    for case in build_testcases.expand(build_testcases.parse(source)):
        elements[f'{case.id} Наименование'] = (case.name, '', False)
        for name, items in case.sections():
            for number, item in enumerate(items, start=1):
                elements[f'{case.id} {build_testcases.LABELS[name]} {number}'] = (item.text, item.expected, bool(item.shared))
    return elements


def check_cases(defects, traps, elements) -> list:
    errors = []
    for row in defects:
        number, case, place = row[0], row[1], row[2]
        found = case_places(case, place)
        if not found and not SECTION.search(place):
            errors.append(f'{number}: не понял место «{place}»')
        for element in sorted(found):
            if elements and element not in elements:
                errors.append(f'{number}: в кейсе {element.split()[0]} нет элемента {element.split(" ", 1)[1]}')
            elif elements and elements[element][2]:
                errors.append(f'{number}: дефект в общем шаге ({element})')
    defect_places = set().union(*(case_places(row[1], row[2]) for row in defects))
    for row in traps:
        found = {p for p in case_places('', row[1]) if not p.startswith(' ')}
        if not found:
            errors.append(f'{row[0]}: не понял место «{row[1]}» — нужен ID кейса, например «900405 Шаг 3»')
        for element in sorted(found):
            if elements and element not in elements:
                errors.append(f'{row[0]}: нет элемента {element}')
        for element in sorted(found & defect_places):
            errors.append(f'{row[0]} и дефект в одном месте: {element}')
    return errors


def protocol_stale(unit: str, source: Path, task: Path) -> bool:
    import build_protocols
    from docx import Document
    docx = task / 'participant' / 'protocols' / f'{source.stem}.docx'
    return not docx.exists() or Document(docx).core_properties.identifier != build_protocols.source_hash(source)


def check_load(defects, traps, task: Path) -> list:
    errors, files, cache = [], load_places.unit_files(task), {}

    def keys_of(number, unit, place):
        kind = load_places.KINDS.get(unit)
        if kind is None:
            errors.append(f'{number}: нет единицы {unit}')
            return set()
        keys = load_places.place_keys(kind, place)
        if not keys:
            errors.append(f'{number}: не понял место «{place}»')
            return set()
        if unit not in files:
            errors.append(f'{number}: нет файла единицы {unit}')
            return keys
        if unit not in cache:
            cache[unit] = load_places.targets(kind, files[unit])
        for key in sorted(keys - cache[unit]):
            errors.append(f'{number}: в единице {unit} нет места {key}')
        return keys

    for row in defects + traps:
        need = CASE_CELLS if row[0].startswith('Д') else 4
        if len(row) - 1 != need:
            errors.append(f'{row[0]}: в строке {len(row) - 1} ячеек, нужно {need} — нет ли «|» в тексте')
    defect_places = set()
    for number, unit, place, category, *_ in defects:
        if number[1:3] != unit:
            errors.append(f'{number}: номер не совпадает с единицей {unit}')
        if category not in CATEGORIES:
            errors.append(f'{number}: категория «{category}», нужна М или 1–5')
        defect_places |= {(unit, key) for key in keys_of(number, unit, place)}
    for number, unit, place, *_ in traps:
        for key in sorted(keys_of(number, unit, place)):
            if (unit, key) in defect_places:
                errors.append(f'{number} и дефект в одном месте: {unit} {key}')
    if (task / 'jury' / 'source' / 'scripts').exists():
        import build_jmx
        errors += [f'сценарий {unit} старше исходника — запустите build_jmx.py' for unit in build_jmx.stale(task)]
    for unit, source in sorted(files.items()):
        if load_places.KINDS[unit] == 'protocol' and protocol_stale(unit, source, task):
            errors.append(f'протокол {unit} старше исходника — запустите build_protocols.py')
    return errors


def check(defects, traps, elements=None, task=None, partial=False):
    errors = []
    per_test = Counter(row[1] for row in defects)
    if not partial:
        if len(per_test) != TESTS:
            errors.append(f'тестов {len(per_test)}, нужно {TESTS}')
        for test, count in per_test.items():
            if count != DEFECTS_PER_TEST:
                errors.append(f'{test}: {count} дефекта, нужно {DEFECTS_PER_TEST}')
        if len(traps) != TRAPS:
            errors.append(f'ловушек {len(traps)}, нужно {TRAPS}')
    if kind_of(defects) == 'load':
        errors += check_load(defects, traps, task)
    elif is_cases(defects):
        errors += check_cases(defects, traps, elements or {})
    else:
        defect_places = set().union(*(places(row[2]) for row in defects)) if defects else set()
        for row in traps:
            for place in sorted(places(row[1]) & defect_places):
                errors.append(f'{row[0]} и дефект в одной строке: {place}')
    if errors:
        sys.exit('Ключ не прошёл проверку:\n' + '\n'.join(errors))


def plain(text: str) -> str:
    """Без разметки Markdown: в Excel обратные кавычки только мешают читать."""
    return text.replace('`', '')


def header(sheet, titles, widths):
    for col, (title, width) in enumerate(zip(titles, widths), start=1):
        cell = sheet.cell(row=1, column=col, value=title)
        cell.font, cell.fill = HEADER, FILL
        sheet.column_dimensions[cell.column_letter].width = width
    sheet.freeze_panes = 'A2'


def build_template(defects, path: Path):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Ответ'
    first = {'cases': ['Кейс', 'Место'], 'load': ['Единица', 'Место']}.get(kind_of(defects), ['Тест', 'Файл:строка'])
    header(sheet, [*first, 'Что не так', 'Почему это проблема', 'Как исправить'], [12, 28, 50, 50, 50])
    tests = list(dict.fromkeys(row[1] for row in defects))
    row = 2
    for test in tests:
        for _ in range(DEFECTS_PER_TEST):
            sheet.cell(row=row, column=1, value=test)
            for col in range(2, 6):
                sheet.cell(row=row, column=col).alignment = WRAP
            row += 1
    book.save(path)


def build_scoring(defects, traps, path: Path):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Подсчёт'
    kind = kind_of(defects)
    cases = kind in ('cases', 'load')
    if cases:
        unit = 'Единица / где' if kind == 'load' else 'Кейс / где'
        header(sheet, ['№', unit, 'Место', 'Категория', 'Суть', 'Засчитываем, если…', '0 или 1'],
               [9, 14, 30, 11, 45, 55, 9])
    else:
        header(sheet, ['№', 'Тест / где', 'Место', 'Суть', 'Засчитываем, если…', '0 или 1'], [9, 14, 30, 45, 55, 9])
    score = 7 if cases else 6
    letter = 'G' if cases else 'F'
    zero_one = DataValidation(type='whole', operator='between', formula1='0', formula2='1', allow_blank=True,
                              showErrorMessage=True, errorTitle='Только 0 или 1',
                              error='Поставьте 1, если засчитываем, или 0, если нет. Другие значения итог не учтёт.')
    sheet.add_data_validation(zero_one)
    row = 2
    for defect in defects:
        for col, value in enumerate(defect, start=1):
            sheet.cell(row=row, column=col, value=plain(value)).alignment = WRAP
        zero_one.add(sheet.cell(row=row, column=score))
        row += 1
    last_defect = row - 1
    row += 1
    what = {'cases': 'кейс верен', 'load': 'место верно'}.get(kind, 'код корректен')
    where = 'место верно' if cases else 'код верен'
    gap = [''] if cases else []
    for col, value in enumerate(['Ловушки', '', '', *gap, f'Почему {what}',
                                 'Ставим 1, если участник предлагает «исправить» это место. '
                                 f'Если сам объяснил, почему {where}, — 0. Каждая ловушка −2 балла'], start=1):
        cell = sheet.cell(row=row, column=col, value=value)
        cell.font, cell.alignment = HEADER, WRAP
    row += 1
    first_trap = row
    for trap in traps:
        number, unit, place, hit, why = trap if kind == 'load' else (trap[0], '', *trap[1:])
        for col, value in enumerate([number, unit, place, *gap, why, hit], start=1):
            sheet.cell(row=row, column=col, value=plain(value)).alignment = WRAP
        zero_one.add(sheet.cell(row=row, column=score))
        row += 1
    last_trap = row - 1
    row += 2
    sheet.cell(row=row, column=score - 1, value='Итог: max(0, дефекты − 2 × ловушки)').font = HEADER
    total = f'=MAX(0,SUM({letter}2:{letter}{last_defect})-2*SUM({letter}{first_trap}:{letter}{last_trap}))'
    sheet.cell(row=row, column=score, value=total).font = HEADER
    book.save(path)


def show_cases(defects, traps, task: Path):
    elements = case_elements(task)
    for row in defects + traps:
        found = case_places(row[1], row[2]) if row[0].startswith('Д') else case_places('', row[1])
        for element in sorted(found):
            text, expected, _ = elements.get(element, ('<нет элемента>', '', False))
            tail = f' | ОР: {expected}' if expected else ''
            print(f'{row[0]} {element}: {text}{tail}'.replace('\n', ' / '))


def show_load(defects, traps, task: Path):
    files = load_places.unit_files(task)
    for row in defects + traps:
        number, unit, place = row[0], row[1], row[2]
        kind = load_places.KINDS.get(unit)
        if kind is None or unit not in files:
            print(f'{number} {unit}: нет файла единицы')
            continue
        for key in sorted(load_places.place_keys(kind, place)):
            print(f'{number} {unit} {key}: {load_places.describe(kind, files[unit], key)}')


def show(defects, traps, task: Path):
    if kind_of(defects) == 'load':
        show_load(defects, traps, task)
        return
    if is_cases(defects):
        show_cases(defects, traps, task)
        return
    skip = {'target', 'jury', '.m2', '.venv', 'build'}
    for row in defects + traps:
        cited = row[2] if row[0].startswith('Д') else row[1]
        for name, start, end in PLACE.findall(cited):
            files = sorted(p for p in task.rglob(name) if not skip & set(p.parts))
            if not files:
                print(f'{row[0]}: файл {name} не найден')
                continue
            for file in files:
                # Одно имя в нескольких песочницах (BaseTest.java) — печатаем каждую с путём.
                label = name if len(files) == 1 else str(file.relative_to(task))
                lines = file.read_text(encoding='utf-8').splitlines()
                for n in range(int(start), int(end or start) + 1):
                    print(f'{row[0]} {label}:{n}: {lines[n - 1].strip() if n <= len(lines) else "<нет строки>"}')


def summary(defects) -> str:
    counts = Counter(row[3] for row in defects)
    return 'Категории: ' + ', '.join(f'{c} — {counts.get(c, 0)}' for c in ['М', '1', '2', '3', '4', '5'])


def build(task: Path, partial=False):
    defects, traps = parse_key(task / 'jury' / 'answer-key.md')
    if kind_of(defects) == 'load':
        check(defects, traps, task=task, partial=partial)
        if not partial:
            build_template(defects, task / 'participant' / 'answer-template.xlsx')
            build_scoring(defects, traps, task / 'jury' / 'scoring-sheet.xlsx')
        print(summary(defects))
        return
    elements = None
    if is_cases(defects):
        export = task / 'participant' / 'testcases.xlsx'
        source = task / 'jury' / 'source' / 'testcases.md'
        if not export.exists():
            sys.exit('нет participant/testcases.xlsx — сначала запустите build_testcases.py')
        elements = case_elements(task)
        if source.exists() and elements != source_elements(source):
            sys.exit('participant/testcases.xlsx старше исходника кейсов — запустите build_testcases.py')
    check(defects, traps, elements)
    build_template(defects, task / 'participant' / 'answer-template.xlsx')
    build_scoring(defects, traps, task / 'jury' / 'scoring-sheet.xlsx')


if __name__ == '__main__':
    if '--help' in sys.argv or len(sys.argv) < 2:
        sys.exit(__doc__)
    target = Path(sys.argv[1]).resolve()
    if '--show' in sys.argv:
        show(*parse_key(target / 'jury' / 'answer-key.md'), target)
    else:
        build(target, partial='--partial' in sys.argv)
        print(f'Готово: {target.name}' + (' (только проверка мест)' if '--partial' in sys.argv else ''))
