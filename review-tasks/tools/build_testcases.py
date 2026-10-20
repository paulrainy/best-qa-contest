"""Выгрузка Test IT для задания 1b из текстового исходника кейсов.

Запуск: python build_testcases.py ../1b-testcase-review [--partial] [--dump ID]
Собирает jury/source/testcases.md → participant/testcases.xlsx
и jury/source/testcases-fixed.md → jury/fixed/testcases.xlsx (если исходник есть).
Падает, если нет кейса 900401–900420, кейсов по BIQ не 7/7/6, в кейсе не 6–12 шагов
или общий шаг не найден. --partial: проверить только кейсы, которые уже написаны.
--dump ID печатает кейс с номерами элементов — по ним пишется место в ключе.

Формат исходника — в docs/superpowers/plans/2026-10-19-1b-testcase-review.md, раздел «Формат исходника».
"""
import copy
import re
import sys
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

COLUMNS = ['ID', 'Расположение', 'Наименование', 'Автоматизирован', '№ шага', 'Предусловия', 'Шаги',
           'Постусловия', 'Ожидаемый результат', 'Тестовые данные', 'Комментарии', 'Итерации', 'Приоритет',
           'Статус', 'Дата создания', 'Автор', 'Тег']
CASE_IDS = [str(n) for n in range(900401, 900421)]
PER_BIQ = {'BIQ-90101': 7, 'BIQ-90102': 7, 'BIQ-90103': 6}
SHARED_LOCATION = 'РТ_ЕФР -> Общие шаги'
SHARED_COMMENT = 'Общий шаг '  # по этой пометке build_sheets.py узнаёт строки общих шагов
SECTIONS = {'Предусловия': 'pre', 'Шаги': 'steps', 'Постусловия': 'post'}
LABELS = {'pre': 'Предусловие', 'steps': 'Шаг', 'post': 'Постусловие'}
WIDE = {'Расположение', 'Наименование', 'Предусловия', 'Шаги', 'Постусловия', 'Ожидаемый результат',
        'Тестовые данные', 'Комментарии'}
HEAD = re.compile(r'^# (\d{6}) (.+)$')
FIELD = re.compile(r'^([А-Яа-яЁё ]+): (.*)$')
REF = re.compile(r'^@(\d{6})$')
WRAP = Alignment(wrap_text=True, vertical='top')


@dataclass
class Item:
    text: str
    expected: str = ''
    data: str = ''
    shared: str | None = None


@dataclass
class Case:
    id: str
    name: str
    fields: dict = field(default_factory=dict)
    pre: list = field(default_factory=list)
    steps: list = field(default_factory=list)
    post: list = field(default_factory=list)

    @property
    def is_shared(self) -> bool:
        return self.fields.get('Расположение') == SHARED_LOCATION

    def sections(self):
        return [(name, getattr(self, name)) for name in ('pre', 'steps', 'post')]


def parse(path: Path) -> list:
    cases, case, section, item, attr = [], None, None, None, None
    for number, line in enumerate(path.read_text(encoding='utf-8').splitlines(), start=1):
        stripped = line.strip()
        if not stripped:
            continue
        head = HEAD.match(line)
        if head:
            case, section, item = Case(head.group(1), head.group(2).strip()), None, None
            cases.append(case)
        elif case is None:
            sys.exit(f'{path.name}, строка {number}: текст до первого кейса')
        elif line.startswith('## ') and stripped[3:] in SECTIONS:
            section, item = getattr(case, SECTIONS[stripped[3:]]), None
        elif section is None and FIELD.match(line):
            key, value = FIELD.match(line).groups()
            case.fields[key] = value.strip()
        elif section is not None and line.startswith('- '):
            item, attr = Item(stripped[2:].strip()), 'text'
            section.append(item)
        elif item is not None and line[0].isspace() and stripped[:2] in ('= ', '> '):
            attr = 'expected' if stripped[0] == '=' else 'data'
            setattr(item, attr, stripped[2:].strip())
        elif item is not None and line[0].isspace():
            setattr(item, attr, getattr(item, attr) + '\n' + stripped)
        else:
            sys.exit(f'{path.name}, строка {number}: не понимаю «{stripped[:60]}»')
    return cases


def biq(case: Case) -> str:
    match = re.search(r'BIQ-\d{5}', case.fields.get('Расположение', ''))
    return match.group(0) if match else ''


def ref(item: Item) -> str | None:
    match = REF.match(item.text)
    return match.group(1) if match else None


def expand(cases: list) -> list:
    shared = {c.id: c for c in cases if c.is_shared}
    result = []
    for case in cases:
        case = copy.deepcopy(case)
        for name, items in case.sections():
            expanded = []
            for item in items:
                record = shared.get(ref(item)) if not case.is_shared else None
                if record:
                    expanded += [Item(s.text, s.expected, s.data, record.id) for s in record.steps]
                else:
                    expanded.append(item)
            setattr(case, name, expanded)
        result.append(case)
    return result


def check(cases: list, ids=CASE_IDS, per_biq=PER_BIQ, steps=(6, 12)) -> None:
    errors = []
    shared = {c.id for c in cases if c.is_shared}
    counts = Counter(c.id for c in cases)
    errors += [f'повтор ID {i}' for i, n in counts.items() if n > 1]
    own = [c for c in cases if not c.is_shared]
    found = {c.id for c in own}
    errors += [f'нет кейса {i}' for i in ids if i not in found]
    errors += [f'лишний кейс {i}' for i in sorted(found - set(ids))]
    biqs = Counter(biq(c) for c in own)
    errors += [f'{b}: кейсов {biqs[b]}, нужно {n}' for b, n in per_biq.items() if biqs[b] != n]
    for case in cases:
        refs = [ref(i) for _, items in case.sections() for i in items if ref(i)]
        if case.is_shared and refs:
            errors.append(f'{case.id}: в общем шаге нельзя подключать другие')
        errors += [f'{case.id}: нет общего шага {r}' for r in refs if r not in shared and not case.is_shared]
    for case in expand(own):
        if not steps[0] <= len(case.steps) <= steps[1]:
            errors.append(f'{case.id}: шагов {len(case.steps)}, нужно {steps[0]}–{steps[1]}')
    if errors:
        sys.exit('Кейсы не прошли проверку:\n' + '\n'.join(errors))


def put(sheet, row: int, values: dict) -> None:
    for name, value in values.items():
        if value:
            cell = sheet.cell(row=row, column=COLUMNS.index(name) + 1, value=value)
            # Текст «= 10 000,00» openpyxl иначе запишет формулой.
            cell.data_type = 's'
            cell.alignment = WRAP


def write(cases: list, path: Path) -> None:
    by_id = {c.id: c for c in cases}
    book = Workbook()
    sheet = book.active
    sheet.title = 'Тест-кейсы'
    for col, name in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=col, value=name)
        cell.font, cell.fill = Font(bold=True), PatternFill('solid', fgColor='DDEBF7')
        sheet.column_dimensions[cell.column_letter].width = 50 if name in WIDE else 13
    sheet.freeze_panes = 'A2'
    row = 2
    for case in [c for c in cases if not c.is_shared] + [c for c in cases if c.is_shared]:
        put(sheet, row, {**case.fields, 'ID': case.id, 'Наименование': case.name})
        row += 1
        for name, items in case.sections():
            column = {'pre': 'Предусловия', 'steps': 'Шаги', 'post': 'Постусловия'}[name]
            for number, item in enumerate(items, start=1):
                comment = f'{SHARED_COMMENT}{item.shared} «{by_id[item.shared].name}»' if item.shared else ''
                put(sheet, row, {'№ шага': f'{LABELS[name]} {number}', column: item.text,
                                 'Ожидаемый результат': item.expected, 'Тестовые данные': item.data,
                                 'Комментарии': comment})
                row += 1
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


def dump(cases: list, case_id: str) -> str:
    case = next(c for c in cases if c.id == case_id)
    lines = [f'{case.id} Наименование: {case.name}']
    for name, items in case.sections():
        for number, item in enumerate(items, start=1):
            label = f'{LABELS[name]} {number}'
            mark = f' [общий шаг {item.shared}]' if item.shared else ''
            lines.append(f'{label}: {item.text}{mark}')
            if item.expected:
                lines.append(f'{label} ОР: {item.expected}')
            if item.data:
                lines.append(f'{label} данные: {item.data}')
    return '\n'.join(lines)


def build(task: Path, partial: bool = False) -> None:
    for name, out in [('testcases.md', task / 'participant' / 'testcases.xlsx'),
                      ('testcases-fixed.md', task / 'jury' / 'fixed' / 'testcases.xlsx')]:
        src = task / 'jury' / 'source' / name
        if not src.exists():
            continue
        cases = parse(src)
        if partial:
            own = [c for c in cases if not c.is_shared]
            ids = [c.id for c in own]
            # Пока кейсы пишутся партиями, число кейсов проверяем только у законченных BIQ.
            done = Counter(biq(c) for c in own)
            check(cases, ids=ids, per_biq={b: n for b, n in PER_BIQ.items() if done[b] >= n})
        else:
            check(cases)
        write(expand(cases), out)
        print(f'Готово: {out.relative_to(task)}')


if __name__ == '__main__':
    if '--help' in sys.argv or len(sys.argv) < 2:
        sys.exit(__doc__)
    target = Path(sys.argv[1]).resolve()
    if '--dump' in sys.argv:
        print(dump(expand(parse(target / 'jury' / 'source' / 'testcases.md')), sys.argv[sys.argv.index('--dump') + 1]))
    else:
        build(target, partial='--partial' in sys.argv)
