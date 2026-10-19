from pathlib import Path

import pytest
from openpyxl import load_workbook

import build_sheets


def make_task(tmp_path: Path, defects_per_test=5, tests=20, traps=10, clash=False) -> Path:
    (tmp_path / 'jury').mkdir()
    (tmp_path / 'participant').mkdir()
    lines = ['# Ключ', '', '| № | Тест | Файл:строки | Суть | Засчитываем, если… |', '| --- | --- | --- | --- | --- |']
    for t in range(tests):
        tk = f'TK9001{t + 1:02d}'
        for d in range(defects_per_test):
            lines.append(f'| Д1{t + 1:02d}.{d + 1} | {tk} | `{tk}Test.java:{10 + d}` | суть | признак |')
    lines += ['', '| № | Где | Попадание | Почему корректен |', '| --- | --- | --- | --- |']
    for n in range(traps):
        place = '`TK900101Test.java:10`' if clash and n == 0 else f'`Helper.java:{100 + n}`'
        lines.append(f'| Л{n + 1} | {place} | замечание | потому что |')
    (tmp_path / 'jury' / 'answer-key.md').write_text('\n'.join(lines), encoding='utf-8')
    return tmp_path


def test_builds_both_sheets(tmp_path):
    task = make_task(tmp_path)
    build_sheets.build(task)

    template = load_workbook(task / 'participant' / 'answer-template.xlsx').active
    assert template['A1'].value == 'Тест'
    assert [template.cell(row=r, column=1).value for r in range(2, 7)] == ['TK900101'] * 5
    assert template.cell(row=101, column=1).value == 'TK900120'

    sheet = load_workbook(task / 'jury' / 'scoring-sheet.xlsx').active
    assert sheet['A2'].value == 'Д101.1'
    assert sheet['A101'].value == 'Д120.5'
    assert sheet['A104'].value == 'Л1'
    assert sheet['A113'].value == 'Л10'
    assert sheet['F116'].value == '=MAX(0,SUM(F2:F101)-2*SUM(F104:F113))'
    validation, = sheet.data_validations.dataValidation
    assert validation.showErrorMessage
    assert validation.error


@pytest.mark.parametrize('kwargs, message', [
    ({'defects_per_test': 4}, 'TK900101: 4 дефекта'),
    ({'traps': 9}, 'ловушек 9'),
    ({'tests': 19}, 'тестов 19'),
    ({'clash': True}, 'TK900101Test.java:10'),
])
def test_rejects_broken_key(tmp_path, kwargs, message):
    task = make_task(tmp_path, **kwargs)
    with pytest.raises(SystemExit, match=message):
        build_sheets.build(task)


# ---------- Задание 1б: место вида «Шаг N» ----------

import build_testcases


def make_case_task(tmp_path: Path, place=None, defects_of_first=5, shared_first=False) -> Path:
    (tmp_path / 'jury').mkdir()
    (tmp_path / 'participant').mkdir()
    lines = ['| № | Кейс | Место | Категория | Суть | Засчитываем, если… |', '| --- | --- | --- | --- | --- | --- |']
    for t in range(1, 21):
        for d in range(1, (defects_of_first if t == 1 else 5) + 1):
            where = place if place and t == 1 and d == 1 else f'Шаг {d}, ожидаемый результат'
            lines.append(f'| Д4{t:02d}.{d} | 9004{t:02d} | {where} | BIQ | суть | признак |')
    lines += ['', '| № | Где | Попадание | Почему верно |', '| --- | --- | --- | --- |']
    lines += [f'| Л{n} | 9004{n:02d} Шаг 6 | замечание | потому что |' for n in range(1, 11)]
    (tmp_path / 'jury' / 'answer-key.md').write_text('\n'.join(lines), encoding='utf-8')
    steps = '\n'.join(f'- Шаг {n}\n  = Результат {n}' for n in range(1, 7))
    source = ['# 900901 Общий шаг\nРасположение: РТ_ЕФР -> Общие шаги\n\n## Шаги\n- Войти\n  = Вошли\n']
    for t in range(1, 21):
        own = steps.replace('- Шаг 1\n  = Результат 1', '- @900901') if shared_first and t == 1 else steps
        source.append(f'# 9004{t:02d} Кейс\nРасположение: BIQ\n\n## Шаги\n{own}\n')
    src = tmp_path / 'testcases.md'
    src.write_text('\n'.join(source), encoding='utf-8')
    cases = build_testcases.expand(build_testcases.parse(src))
    build_testcases.write(cases, tmp_path / 'participant' / 'testcases.xlsx')
    return tmp_path


def test_builds_case_sheets(tmp_path):
    task = make_case_task(tmp_path)
    build_sheets.build(task)

    template = load_workbook(task / 'participant' / 'answer-template.xlsx').active
    assert [template['A1'].value, template['B1'].value, template['A2'].value] == ['Кейс', 'Место', '900401']
    assert template.cell(row=101, column=1).value == '900420'

    sheet = load_workbook(task / 'jury' / 'scoring-sheet.xlsx').active
    assert sheet['D1'].value == 'Категория'
    assert sheet['A2'].value == 'Д401.1'
    assert sheet['D2'].value == 'BIQ'
    assert sheet['A104'].value == 'Л1'
    assert sheet['G116'].value == '=MAX(0,SUM(G2:G101)-2*SUM(G104:G113))'


@pytest.mark.parametrize('kwargs, message', [
    ({'place': 'Шаг 6'}, 'Л1 и дефект в одном месте: 900401 Шаг 6'),
    ({'place': 'Шаг 9'}, 'Д401.1: в кейсе 900401 нет элемента Шаг 9'),
    ({'place': 'где-то в середине'}, 'Д401.1: не понял место «где-то в середине»'),
    ({'shared_first': True}, r'Д401.1: дефект в общем шаге \(900401 Шаг 1\)'),
    ({'defects_of_first': 4}, '900401: 4 дефекта'),
])
def test_rejects_broken_case_key(tmp_path, kwargs, message):
    task = make_case_task(tmp_path, **kwargs)
    with pytest.raises(SystemExit, match=message):
        build_sheets.build(task)


def test_whole_section_place_is_allowed(tmp_path):
    build_sheets.build(make_case_task(tmp_path, place='Предусловия — нет клиента'))


def test_place_may_name_other_case():
    assert build_sheets.case_places('900401', 'Шаг 2; 900410 Предусловие 1') == {'900401 Шаг 2', '900410 Предусловие 1'}


def test_show_prints_case_element(tmp_path, capsys):
    task = make_case_task(tmp_path)
    build_sheets.show(*build_sheets.parse_key(task / 'jury' / 'answer-key.md'), task)
    out = capsys.readouterr().out
    assert 'Д401.1 900401 Шаг 1: Шаг 1 | ОР: Результат 1' in out
    assert 'Л1 900401 Шаг 6: Шаг 6 | ОР: Результат 6' in out


def test_case_key_needs_export(tmp_path):
    task = make_case_task(tmp_path)
    (task / 'participant' / 'testcases.xlsx').unlink()
    with pytest.raises(SystemExit, match='нет participant/testcases.xlsx'):
        build_sheets.build(task)


def test_case_key_accepts_built_export(tmp_path):
    task = make_case_task(tmp_path)
    (task / 'jury' / 'source').mkdir()
    (tmp_path / 'testcases.md').rename(task / 'jury' / 'source' / 'testcases.md')
    build_sheets.build(task)


def test_case_key_rejects_stale_export(tmp_path):
    task = make_case_task(tmp_path)
    (task / 'jury' / 'source').mkdir()
    source = task / 'jury' / 'source' / 'testcases.md'
    edited = (tmp_path / 'testcases.md').read_text(encoding='utf-8').replace('- Шаг 3\n', '- Шаг 3 исправлен\n', 1)
    source.write_text(edited, encoding='utf-8')
    with pytest.raises(SystemExit, match='testcases.xlsx старше исходника'):
        build_sheets.build(task)


# ---------- Задание 1г: сценарии, протоколы, дашборды, SQL ----------

import json

import jmx
import load_places


def write_jmx(path: Path):
    samplers = [jmx.http_request(f'S{n}', 'GET', f'/s{n}') for n in range(1, 7)]
    plan = jmx.test_plan('x', [jmx.thread_group('Thread Group', 1, 1, children=samplers)])
    path.write_text(jmx.render(plan), encoding='utf-8')


def write_dashboard(path: Path):
    panels = [{'id': n, 'title': f'P{n}', 'type': 'timeseries', 'gridPos': {'x': 0, 'y': n, 'w': 12, 'h': 8},
               'datasource': {'type': 'influxdb', 'uid': 'influxdb-nt'}, 'targets': [{'refId': 'A', 'query': 'q'}]}
              for n in range(1, 7)]
    data = {'title': 'D', 'uid': path.stem, 'schemaVersion': 39, 'panels': panels, 'templating': {'list': []}}
    path.write_text(json.dumps(data), encoding='utf-8')


LOAD_PLACE = {'jmx': '`Thread Group → S{n}`', 'sql': 'строка {n}', 'json': 'Панель «P{n}»'}


def make_load_task(tmp_path: Path, monkeypatch, place=None, category=None, number=None, defects_of_first=5,
                   trap_place=None, units=range(1, 21), traps=10) -> Path:
    # До сборщика протоколов единицы 11–15 в фикстуре — SQL
    for unit in range(11, 16):
        monkeypatch.setitem(load_places.KINDS, f'{unit:02d}', 'sql')
    (tmp_path / 'jury').mkdir()
    for folder in ('scripts', 'monitoring'):
        (tmp_path / 'participant' / folder).mkdir(parents=True)
    lines = ['| № | Единица | Место | Категория | Суть | Засчитываем, если… |', '| --- | --- | --- | --- | --- | --- |']
    categories = ['М', '1', '2', '3', '4']
    for u in units:
        unit = f'{u:02d}'
        kind = load_places.KINDS[unit]
        if kind == 'jmx':
            write_jmx(tmp_path / 'participant' / 'scripts' / f'{unit}-s.jmx')
        elif kind == 'json':
            write_dashboard(tmp_path / 'participant' / 'monitoring' / f'{unit}-d.json')
        else:
            (tmp_path / 'participant' / 'monitoring' / f'{unit}-q.sql').write_text(
                ''.join(f'SELECT {n};\n' for n in range(1, 9)), encoding='utf-8')
        for d in range(1, (defects_of_first if u == 1 else 5) + 1):
            first = u == 1 and d == 1
            where = place if place and first else LOAD_PLACE[kind].format(n=d)
            cat = category if category and first else categories[d - 1]
            num = number if number and first else f'Д{unit}.{d}'
            lines.append(f'| {num} | {unit} | {where} | {cat} | суть | признак |')
    lines += ['', '| № | Единица | Где | Попадание | Почему верно |', '| --- | --- | --- | --- | --- |']
    for n in range(1, traps + 1):
        where = trap_place if trap_place and n == 1 else '`Thread Group → S6`'
        lines.append(f'| Л{n} | {n:02d} | {where} | замечание | потому что |')
    (tmp_path / 'jury' / 'answer-key.md').write_text('\n'.join(lines), encoding='utf-8')
    return tmp_path


def test_builds_load_sheets(tmp_path, monkeypatch):
    task = make_load_task(tmp_path, monkeypatch)
    build_sheets.build(task)

    template = load_workbook(task / 'participant' / 'answer-template.xlsx').active
    assert [template['A1'].value, template['B1'].value, template['A2'].value] == ['Единица', 'Место', '01']
    assert template.cell(row=101, column=1).value == '20'

    sheet = load_workbook(task / 'jury' / 'scoring-sheet.xlsx').active
    assert sheet['B1'].value == 'Единица / где'
    assert sheet['D1'].value == 'Категория'
    assert sheet['A2'].value == 'Д01.1'
    assert sheet['A104'].value == 'Л1'
    assert sheet['B104'].value == '01'
    assert sheet['G116'].value == '=MAX(0,SUM(G2:G101)-2*SUM(G104:G113))'


@pytest.mark.parametrize('kwargs, message', [
    ({'place': '`Thread Group → S9`'}, 'Д01.1: в единице 01 нет места jmx:Thread Group → S9'),
    ({'trap_place': '`Thread Group → S1`'}, 'Л1 и дефект в одном месте: 01 jmx:Thread Group → S1'),
    ({'category': '7'}, 'Д01.1: категория «7»'),
    ({'number': 'Д02.1'}, 'Д02.1: номер не совпадает с единицей 01'),
    ({'place': 'где-то'}, 'Д01.1: не понял место «где-то»'),
    ({'defects_of_first': 4}, '01: 4 дефекта'),
])
def test_rejects_broken_load_key(tmp_path, monkeypatch, kwargs, message):
    task = make_load_task(tmp_path, monkeypatch, **kwargs)
    with pytest.raises(SystemExit, match=message):
        build_sheets.build(task)


def test_partial_skips_counts(tmp_path, monkeypatch):
    task = make_load_task(tmp_path, monkeypatch, units=[1], traps=1)
    build_sheets.build(task, partial=True)
    assert not (task / 'participant' / 'answer-template.xlsx').exists()


def test_partial_still_checks_places(tmp_path, monkeypatch):
    task = make_load_task(tmp_path, monkeypatch, units=[1], traps=1, place='`Thread Group → S9`')
    with pytest.raises(SystemExit, match='нет места jmx:Thread Group → S9'):
        build_sheets.build(task, partial=True)


SCENARIO = '''
import jmx

FILE = '01-s'


def plan(fixed):
    samplers = [jmx.http_request(f'S{n}', 'GET', f'/s{n}') for n in range(1, 7)]
    return jmx.test_plan('x', [jmx.thread_group('Thread Group', 1, 1, children=samplers)])
'''


def test_rejects_stale_jmx(tmp_path, monkeypatch):
    import build_jmx
    task = make_load_task(tmp_path, monkeypatch)
    (task / 'jury' / 'source' / 'scripts').mkdir(parents=True)
    (task / 'jury' / 'source' / 'scripts' / '01_s.py').write_text(SCENARIO, encoding='utf-8')
    build_jmx.build(task)
    build_sheets.build(task)
    path = task / 'participant' / 'scripts' / '01-s.jmx'
    path.write_text(path.read_text(encoding='utf-8').replace('/s1<', '/s1x<'), encoding='utf-8')
    with pytest.raises(SystemExit, match='сценарий 01 старше исходника — запустите build_jmx.py'):
        build_sheets.build(task)


def test_show_prints_load_places(tmp_path, monkeypatch, capsys):
    task = make_load_task(tmp_path, monkeypatch)
    build_sheets.show(*build_sheets.parse_key(task / 'jury' / 'answer-key.md'), task)
    out = capsys.readouterr().out
    assert 'Д01.1 01 jmx:Thread Group → S1: HTTPSamplerProxy' in out
    assert 'HTTPSampler.path=/s1' in out
    assert 'Д16.2 16 панель:P2: timeseries' in out
    assert 'Д19.3 19 строка:3: SELECT 3;' in out
    assert 'Л1 01 jmx:Thread Group → S6' in out
